"""选股结果导出 Excel（P4-S3）。

职责：
    * 把 :class:`~app.engine.screener.ScreenResult` 导出为 ``.xlsx``，列 = 代码 + 名称 +
      各选股条件的指标实际值；
    * **纯逻辑与写盘分离**：:func:`build_workbook` 只构造 ``openpyxl.Workbook``（不碰磁盘），
      :func:`export_screen_result` 负责落盘并把一切失败归一为 :class:`ExcelExportError`；
    * 向后兼容：``metrics`` 为空（旧调用方）时降级为「代码 + 名称」两列，不报错。

列头约定：
    * 前两列固定「代码」「名称」；
    * 之后每个条件一列，列头默认取条件字段名（如 ``pe``、``close``、
      ``close.rolling(20).mean()``）；同名字段自动追加 ``op value`` 消歧，
      保证「列头与条件一一对应、可读」。

数值格式：
    * 指标值一律按 ``float`` 写盘（``None`` 留空），文件可被 Excel / openpyxl 正常打开。
"""
from __future__ import annotations

import logging
import math
from collections import Counter
from pathlib import Path
from typing import Any, Callable, Optional, Sequence, Union

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# 固定表头。
_CODE_HEADER: str = "代码"
_NAME_HEADER: str = "名称"


class ExcelExportError(Exception):
    """导出失败统一异常（由 UI 层捕获并提示用户）。"""


# --------------------------------------------------------------------------- #
# 列头推导（纯函数）
# --------------------------------------------------------------------------- #
def _condition_field(cond: Any) -> str:
    """从 Condition / dict 中提取字段名（非预期形态返回空串）。"""
    if isinstance(cond, dict):
        return str(cond.get("field", "")).strip()
    return str(getattr(cond, "field", "")).strip()


def _condition_headers(
    conditions: Optional[Sequence[Any]], metrics: Sequence[dict]
) -> list[str]:
    """推导指标列头，保证与条件一一对应且可读。

    优先使用 ``conditions``（含 ``op/value``，可对同名字段消歧）；``conditions`` 为
    ``None`` 时从 ``metrics[0]["conditions"]`` 的字段名推导。两者皆空返回空列表。
    """
    if conditions:
        fields = [_condition_field(c) for c in conditions]
        counts = Counter(fields)
        raw: list[str] = []
        for i, cond in enumerate(conditions):
            field = fields[i]
            if counts[field] > 1:
                # 同名字段（如两个 ``close`` 条件）→ 追加操作符与阈值以消歧。
                if isinstance(cond, dict):
                    op = str(cond.get("op", "")).strip()
                    value = cond.get("value", "")
                else:
                    op = str(getattr(cond, "op", "")).strip()
                    value = getattr(cond, "value", "")
                raw.append(f"{field} {op} {value}".strip())
            else:
                raw.append(field)

        # 兜底：完全相同的条件（field+op+value 全同）经上面消歧后仍会撞名，
        # 与 metrics 派生路径风格一致，对最终仍重复的列头统一追加 (n) 编号，保证唯一可读。
        header_counts = Counter(raw)
        seen: Counter = Counter()
        headers = []
        for header in raw:
            if header_counts[header] > 1:
                seen[header] += 1
                headers.append(f"{header} ({seen[header]})")
            else:
                headers.append(header)
        return headers

    if not metrics:
        return []
    first = metrics[0] or {}
    items = list((first.get("conditions") or []))
    fields = [str(item.get("field", "")).strip() if isinstance(item, dict) else "" for item in items]
    counts = Counter(fields)
    seen: Counter = Counter()
    derived: list[str] = []
    for field in fields:
        if counts[field] > 1:
            seen[field] += 1
            derived.append(f"{field}({seen[field]})")
        else:
            derived.append(field)
    return derived


# --------------------------------------------------------------------------- #
# 数值归一
# --------------------------------------------------------------------------- #
def _to_cell_value(value: Any) -> Any:
    """把指标值归一为可写盘的值：``None`` 留空，数值转 ``float``，其它转 ``str``。

    引擎约定指标值为 ``float | None``；此处对异常类型做兜底，避免写盘报错。
    """
    if value is None:
        return None
    if isinstance(value, bool):
        return float(value)
    try:
        num = float(value)
    except (TypeError, ValueError):
        return str(value)
    if not math.isfinite(num):
        return None
    return num


def _resolve_name(
    symbol: str, name_resolver: Optional[Callable[[str], Optional[str]]]
) -> Optional[str]:
    """尽力解析代码 → 名称；取不到名 / 解析异常返回 ``None``（写盘留空，不崩溃）。"""
    if name_resolver is None:
        return None
    try:
        name = name_resolver(symbol)
    except Exception as exc:  # noqa: BLE001 - 名称解析失败不影响导出
        logger.debug("name_resolver(%s) 失败：%s", symbol, exc)
        return None
    return str(name) if name else None


# --------------------------------------------------------------------------- #
# 纯逻辑：构造工作簿
# --------------------------------------------------------------------------- #
def build_workbook(
    result: Any,
    name_resolver: Optional[Callable[[str], Optional[str]]] = None,
    conditions: Optional[Sequence[Any]] = None,
) -> Workbook:
    """构造导出工作簿（纯逻辑，不写盘）。

    Args:
        result: :class:`~app.engine.screener.ScreenResult`（或具备 ``hits`` / ``metrics``
            属性的等价对象）。
        name_resolver: 可选「代码 → 名称」回调。
        conditions: 可选条件列表（``Condition`` 或 dict），用于生成指标列头；为 ``None``
            时从 ``result.metrics[0]`` 的 ``conditions`` 字段名推导。

    Returns:
        ``openpyxl.Workbook``：首个工作表含表头与数据行。
    """
    hits = list(getattr(result, "hits", []) or [])
    metrics = list(getattr(result, "metrics", []) or [])
    headers = [_CODE_HEADER, _NAME_HEADER] + _condition_headers(conditions, metrics)
    metric_column_count = len(headers) - 2

    wb = Workbook()
    ws = wb.active
    ws.title = "选股结果"

    # 表头（加粗）。
    bold = Font(bold=True)
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = bold

    # 数据行：每个命中股一行。
    for row_index, symbol in enumerate(hits):
        code = str(symbol)
        row = row_index + 2
        ws.cell(row=row, column=1, value=code)
        ws.cell(row=row, column=2, value=_resolve_name(code, name_resolver))

        entry = metrics[row_index] if row_index < len(metrics) else None
        metric_items = []
        if isinstance(entry, dict):
            metric_items = list(entry.get("conditions") or [])
        for col_index in range(metric_column_count):
            value = None
            if col_index < len(metric_items):
                item = metric_items[col_index]
                if isinstance(item, dict):
                    value = item.get("value")
            ws.cell(row=row, column=3 + col_index, value=_to_cell_value(value))

    # 列宽自适应（简单启发式：按表头长度 + 少许余量），不强制，保证可读。
    for col, header in enumerate(headers, start=1):
        width = max(len(str(header)) * 2 + 2, 10)
        ws.column_dimensions[get_column_letter(col)].width = min(width, 40)

    return wb


# --------------------------------------------------------------------------- #
# 写盘入口
# --------------------------------------------------------------------------- #
def export_screen_result(
    result: Any,
    path: Union[str, Path],
    name_resolver: Optional[Callable[[str], Optional[str]]] = None,
    conditions: Optional[Sequence[Any]] = None,
) -> str:
    """把选股结果导出为 ``.xlsx`` 文件。

    Args:
        result: 选股结果（``ScreenResult``）。
        path: 目标文件路径；扩展名缺失时自动补 ``.xlsx``。
        name_resolver: 可选「代码 → 名称」回调。
        conditions: 可选条件列表（用于生成指标列头）。

    Returns:
        str: 实际写入的文件路径（``Path`` 转 ``str``）。

    Raises:
        ExcelExportError: 目录不存在 / 路径非法 / 写盘失败，异常信息含可读原因。
    """
    try:
        target = Path(path)
        if not target.parent.exists():
            raise ExcelExportError(f"保存目录不存在：{target.parent}")
        if target.suffix.lower() != ".xlsx":
            target = target.with_suffix(".xlsx")
        wb = build_workbook(result, name_resolver=name_resolver, conditions=conditions)
        wb.save(target)
    except ExcelExportError:
        raise
    except Exception as exc:  # noqa: BLE001 - 统一归一为明确异常，绝不静默吞掉
        logger.exception("导出选股结果失败")
        raise ExcelExportError(f"导出失败：{exc}") from exc
    return str(target)


__all__ = [
    "ExcelExportError",
    "build_workbook",
    "export_screen_result",
]
