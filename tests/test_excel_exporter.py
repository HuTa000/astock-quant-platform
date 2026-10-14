"""P4-S3 选股结果导出 Excel —— 纯导出逻辑测试。

覆盖：
    1. 有 metrics → 导出「代码 / 名称 / 指标列」齐全、数值正确（读回 xlsx 断言）；
    2. metrics 为空 → 降级为「代码 + 名称」两列，不报错；
    3. name_resolver 取不到名 / 抛异常 → 不崩溃、名称为空；
    4. 指标值为 None → 单元格留空；
    5. 写盘失败（目录不存在）→ 抛 :class:`ExcelExportError`，不静默；
    6. 列头推导：``conditions`` 提供时用字段名、同名字段消歧；为 None 时从 metrics 派生；
    7. ``build_workbook`` 纯逻辑不写盘；``ScreenResult`` 不带 metrics 向后兼容。
"""
from __future__ import annotations

import openpyxl
import pytest

from app.engine.screener import Condition, ScreenResult
from app.export.excel_exporter import (
    ExcelExportError,
    build_workbook,
    export_screen_result,
)


# --------------------------------------------------------------------------- #
# 辅助
# --------------------------------------------------------------------------- #
def _result(hits, metrics=None) -> ScreenResult:
    return ScreenResult(hits=list(hits), failed=[], total=len(hits), metrics=list(metrics or []))


def _read(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    return list(ws.iter_rows(values_only=True))


def _headers(rows) -> list:
    return list(rows[0]) if rows else []


# --------------------------------------------------------------------------- #
# 1. 有 metrics → 三部分齐全、数值正确
# --------------------------------------------------------------------------- #
def test_export_with_metrics_writes_code_name_and_values(tmp_path) -> None:
    result = _result(
        ["000001", "600519"],
        metrics=[
            {
                "symbol": "000001",
                "conditions": [
                    {"field": "pe", "value": 12.5},
                    {"field": "close", "value": 10.0},
                ],
            },
            {
                "symbol": "600519",
                "conditions": [
                    {"field": "pe", "value": 30.0},
                    {"field": "close", "value": None},
                ],
            },
        ],
    )
    path = tmp_path / "result.xlsx"

    saved = export_screen_result(
        result,
        path,
        name_resolver=lambda s: {"000001": "平安银行", "600519": "贵州茅台"}.get(s),
        conditions=[Condition("pe", "<", 20), Condition("close", ">", 5)],
    )

    assert saved == str(path)
    rows = _read(path)
    assert _headers(rows) == ["代码", "名称", "pe", "close"]
    assert rows[1][0] == "000001"
    assert rows[1][1] == "平安银行"
    assert rows[1][2] == 12.5
    assert rows[1][3] == 10.0
    assert rows[2][0] == "600519"
    assert rows[2][1] == "贵州茅台"
    assert rows[2][2] == 30.0
    assert rows[2][3] is None  # None 指标值 → 留空


# --------------------------------------------------------------------------- #
# 2. metrics 为空 → 降级两列
# --------------------------------------------------------------------------- #
def test_export_empty_metrics_falls_back_two_columns(tmp_path) -> None:
    result = _result(["000001"], metrics=[])
    path = tmp_path / "fallback.xlsx"

    export_screen_result(result, path)

    rows = _read(path)
    assert _headers(rows) == ["代码", "名称"]
    assert rows[1][0] == "000001"
    assert rows[1][1] is None  # 无 name_resolver → 名称留空


# --------------------------------------------------------------------------- #
# 3. name_resolver 取不到名 / 抛异常 → 不崩溃
# --------------------------------------------------------------------------- #
def test_export_name_resolver_missing_name_no_crash(tmp_path) -> None:
    result = _result(["000001"], metrics=[{"symbol": "000001", "conditions": []}])
    path = tmp_path / "name.xlsx"

    export_screen_result(result, path, name_resolver=lambda s: None)

    rows = _read(path)
    assert rows[1][1] is None


def test_export_name_resolver_raises_no_crash(tmp_path) -> None:
    result = _result(["000001"], metrics=[{"symbol": "000001", "conditions": []}])
    path = tmp_path / "name2.xlsx"

    def bad_resolver(symbol: str):
        raise RuntimeError("resolve boom")

    export_screen_result(result, path, name_resolver=bad_resolver)

    rows = _read(path)
    assert rows[1][1] is None


# --------------------------------------------------------------------------- #
# 4. None 指标值留空（单列显式断言）
# --------------------------------------------------------------------------- #
def test_export_none_metric_value_leaves_cell_empty(tmp_path) -> None:
    result = _result(
        ["000001"],
        metrics=[{"symbol": "000001", "conditions": [{"field": "roe", "value": None}]}],
    )
    path = tmp_path / "none.xlsx"

    export_screen_result(result, path, conditions=[Condition("roe", ">", 0)])

    rows = _read(path)
    assert _headers(rows) == ["代码", "名称", "roe"]
    assert rows[1][2] is None


# --------------------------------------------------------------------------- #
# 5. 写盘失败（目录不存在）→ 抛明确异常，不静默
# --------------------------------------------------------------------------- #
def test_export_invalid_directory_raises(tmp_path) -> None:
    result = _result(["000001"], metrics=[])
    bad_path = tmp_path / "no_such_dir" / "x.xlsx"

    with pytest.raises(ExcelExportError):
        export_screen_result(result, bad_path)


# --------------------------------------------------------------------------- #
# 6. 列头推导
# --------------------------------------------------------------------------- #
def test_headers_derived_from_metrics_when_conditions_none(tmp_path) -> None:
    result = _result(
        ["000001"],
        metrics=[
            {
                "symbol": "000001",
                "conditions": [
                    {"field": "pe", "value": 1.0},
                    {"field": "close.rolling(20).mean()", "value": 2.0},
                ],
            }
        ],
    )
    path = tmp_path / "derived.xlsx"

    export_screen_result(result, path)

    rows = _read(path)
    assert _headers(rows) == ["代码", "名称", "pe", "close.rolling(20).mean()"]


def test_duplicate_condition_fields_disambiguated(tmp_path) -> None:
    result = _result(
        ["000001"],
        metrics=[
            {
                "symbol": "000001",
                "conditions": [
                    {"field": "close", "value": 6.0},
                    {"field": "close", "value": 6.0},
                ],
            }
        ],
    )
    path = tmp_path / "dup.xlsx"

    export_screen_result(
        result,
        path,
        conditions=[Condition("close", ">", 0), Condition("close", "<", 5)],
    )

    rows = _read(path)
    assert _headers(rows) == ["代码", "名称", "close > 0.0", "close < 5.0"]


def test_identical_condition_headers_numbered(tmp_path) -> None:
    """两个完全相同的条件（field+op+value 全同）→ 列头唯一，且与数据列一一对应。"""
    result = _result(
        ["000001"],
        metrics=[
            {
                "symbol": "000001",
                "conditions": [
                    {"field": "close", "value": 5.0},
                    {"field": "close", "value": 6.0},
                ],
            }
        ],
    )
    path = tmp_path / "identical.xlsx"

    export_screen_result(
        result,
        path,
        conditions=[Condition("close", ">", 5), Condition("close", ">", 5)],
    )

    rows = _read(path)
    assert _headers(rows) == ["代码", "名称", "close > 5.0 (1)", "close > 5.0 (2)"]
    # 数据列仍与条件一一对应：两列分别取到 5.0 / 6.0
    assert rows[1][2] == 5.0
    assert rows[1][3] == 6.0


# --------------------------------------------------------------------------- #
# 7. 纯逻辑不写盘 + 向后兼容
# --------------------------------------------------------------------------- #
def test_build_workbook_is_pure_and_returns_workbook() -> None:
    result = _result(["000001"], metrics=[])
    wb = build_workbook(result, name_resolver=lambda s: "平安银行")
    ws = wb.active
    assert ws["A1"].value == "代码"
    assert ws["B1"].value == "名称"


def test_screen_result_backward_compat_default_metrics_empty() -> None:
    result = ScreenResult(hits=["000001"], failed=[], total=1)
    assert result.metrics == []
    assert result.hits == ["000001"]
    assert result.hit_count == 1
